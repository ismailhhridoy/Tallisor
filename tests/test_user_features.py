"""Features a real user needs: settings, saved history, Excel report, auto-approval."""
import io
import json

import pytest
from openpyxl import load_workbook

from tallisor import settings, store
from tallisor.report import build_workbook
from tallisor.security.policy import PermissionDenied
from tallisor.session import Session

REAL = __import__("pathlib").Path(__file__).resolve().parents[1] / "sample_data" / "realformats"


# ------------------------------------------------------------------ settings
def test_settings_round_trip_and_validation():
    p = settings.default_profile("Kestrel Goods")
    p["fee_schedules"]["Stripe"] = settings.fee_schedule(2.9, 0.30, 4.4, 0.30, "us")
    p["cash"]["payments"] = settings.payments_from_rows([
        {"name": "Rent", "amount": 1500, "repeats": "monthly", "next_date": "day 1"},
        {"name": "Payroll", "amount": 3800, "repeats": "every 2 weeks", "next_date": "2026-10-23"},
        {"name": "Tax", "amount": 900.5, "repeats": "once", "next_date": "2026-11-30"},
        {"name": "", "amount": 5, "repeats": "weekly", "next_date": "2026-10-20"},   # blank row ignored
    ])
    settings.save(p)
    back = settings.load()
    assert back["fee_schedules"]["Stripe"] == {"domestic": {"bp": 290, "fixed": 30},
                                               "international": {"bp": 440, "fixed": 30}, "home_country": "US"}
    assert back["cash"]["payments"] == [
        {"name": "Rent", "amount": 150000, "monthly_day": 1},
        {"name": "Payroll", "amount": 380000, "every_days": 14, "first": "2026-10-23"},
        {"name": "Tax", "amount": 90050, "once": "2026-11-30"}]
    assert settings.payment_rows(back)[1] == {"name": "Payroll", "amount": 3800.0, "repeats": "every 2 weeks",
                                               "next_date": "2026-10-23"}
    bad = settings.default_profile("x")
    bad["fee_schedules"]["Odd"] = settings.fee_schedule(35, 0)
    bad["cash"]["floor"] = -1
    with pytest.raises(ValueError, match="between 0% and 20%.*Safety floor"):
        settings.validate(bad)
    with pytest.raises(ValueError, match="next date"):
        settings.payments_from_rows([{"name": "Loan", "amount": 10, "repeats": "weekly", "next_date": ""}])


def test_missing_fee_contract_does_not_invent_overcharges():
    profile = json.loads((REAL / "profile.json").read_text())
    del profile["fee_schedules"]["PayPal"]
    s = Session.from_dir(REAL).__class__(**{**_real_sources(), "profile": profile}).run()
    assert settings.missing_fee_contracts(profile, s.recon.providers) == ["PayPal"]
    assert not [c for c in s.recon.cases if c.category == "fee_variance"]
    assert s.recon.residual == 0


def _real_sources():
    f = json.loads((REAL / "profile.json").read_text())["files"]
    return {"books": REAL / f["books"], "processors": [REAL / p for p in f["processors"]], "bank": REAL / f["bank"]}


# ------------------------------------------------------------------ saved runs
def test_saved_run_reopens_with_decisions_replayed(session):
    c = session.recon.case("NW-48241")
    session.approve(c.id, role="owner", proposal_id=c.proposal.id, person="Maya Chen")
    session.escalate("NW-48245", role="reviewer", person="Sam", note="ask the bank")
    run_id = store.save_run(session)
    assert store.list_runs()[0]["id"] == run_id and store.list_runs()[0]["approved"] == 1

    again, report = store.open_run(run_id)
    assert report == {"applied": 2, "skipped": []}
    assert again.recon.case("NW-48241").status == "approved"
    assert again.recon.case("NW-48245").status == "escalated"
    assert again.recon.gap == session.recon.gap
    replayed = [e for e in again.log.events if e["details"].get("replayed")]
    assert replayed and replayed[0]["details"]["person"] == "Maya Chen"


def test_replay_skips_a_decision_whose_fix_changed(session):
    c = session.recon.case("NW-48242")
    report = session.__class__.from_dir(session.sources["books"].parent).run().replay(
        [{"action": "approve_fix", "ref": c.ref, "proposal": "0000000000000000", "person": "x", "role": "owner"}])
    assert report["applied"] == 0 and "changed" in report["skipped"][0]["reason"]


# ------------------------------------------------------------------ Excel report
def test_excel_report_has_every_sheet_and_totals_match(session):
    wb = load_workbook(io.BytesIO(build_workbook(session)))
    assert wb.sheetnames == ["Summary", "Differences", "Journal", "Cash forecast", "Alerts", "Financing",
                             "Audit log"]
    diffs = wb["Differences"]
    impacts = [row[5].value for row in diffs.iter_rows(min_row=2) if row[0].value not in ("Total",)]
    assert round(sum(impacts) * 100) == session.recon.gap
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=3)}
    assert round(summary["Gap"] * 100) == session.recon.gap
    assert summary["Audit log head hash"] == session.log.head
    journal = list(wb["Journal"].iter_rows(min_row=2, values_only=True))[:-1]
    assert round(sum(r[5] or 0 for r in journal), 2) == round(sum(r[6] or 0 for r in journal), 2)
    assert wb["Audit log"].max_row == len(session.log.events) + 1


# ------------------------------------------------------------------ auto-approval rules
def _with_rule(session, **rule):
    session.profile = {**session.profile, "auto_approve": {"enabled": True, "set_by": "Maya Chen", **rule}}
    return session


def test_auto_approval_is_off_by_default(session):
    assert session.apply_rules(person="Maya Chen") == []


def test_auto_approval_only_small_high_confidence_matching_categories(session):
    _with_rule(session, categories=["fee_variance", "fx_difference", "payment_failed"], max_amount=2000)
    done = session.apply_rules(person="Maya Chen")
    refs = {session.recon.case(i).ref for i in done}
    assert refs == {"NW-48240", "NW-48243"}          # fee 18.60 and FX 16.51; failed payment is medium confidence
    ev = [e for e in session.log.events if e["action"] == "approve_fix"]
    assert all(e["actor"] == "owner" and "set by Maya Chen" in e["details"]["rule"] for e in ev)


def test_auto_approval_stops_with_kill_switch_and_needs_owner(session):
    _with_rule(session, categories=["fee_variance"], max_amount=10_000)
    with pytest.raises(PermissionDenied):
        session.apply_rules(person="Sam", role="reviewer")
    session.halt("owner", "Maya Chen")
    assert session.apply_rules(person="Maya Chen") == []
    assert session.recon.case("NW-48240").status == "open"


def test_same_files_find_the_saved_run(session):
    run_id = store.save_run(session)
    files = {"books": session.sources["books"], "processors": session.sources["processors"],
             "bank": session.sources["bank"], "history": session.history_src,
             "receivables": session.receivables_src}
    assert store.find_by_inputs(files) == run_id
    assert store.find_by_inputs({**files, "receivables": None}) is None


def test_reopen_with_new_settings_skips_decisions_whose_fix_changed(session):
    c = session.recon.case("NW-48240")                       # fee difference, depends on the fee contract
    session.approve(c.id, role="owner", proposal_id=c.proposal.id, person="Maya Chen")
    r = session.recon.case("NW-48241")                       # unrecorded refund, independent of fees
    session.approve(r.id, role="owner", proposal_id=r.proposal.id, person="Maya Chen")
    run_id = store.save_run(session)
    profile = json.loads(json.dumps(session.profile))
    profile["fee_schedules"]["PayRail"]["international"] = {"bp": 400, "fixed": 30}
    again, report = store.open_run(run_id, profile=profile)
    assert report["applied"] == 1 and again.recon.case("NW-48241").status == "approved"
    assert [d["ref"] for d in report["skipped"]] == ["NW-48240"]
    historical, _ = store.open_run(run_id)
    assert historical.recon.case("NW-48240").status == "approved"
